# Test template: a variable split across runs (as Word often does), plain variables, and a table row repeated per item.
from docx import Document
d = Document()
d.add_heading('Devis {{numero}}', 1)
p = d.add_paragraph('Bonjour ')
p.add_run('{{client_')          # split variable: "{{client_" + "nom}}" in two runs, the second one bold
r = p.add_run('nom}}'); r.bold = True
p.add_run(', voici notre proposition du {{date}}.')
t = d.add_table(rows=2, cols=3); t.style = 'Table Grid'
for i, h in enumerate(['Désignation', 'Quantité', 'Prix']): t.rows[0].cells[i].text = h
for i, v in enumerate(['{{#lignes}}{{designation}}', '{{quantite}}', '{{prix}}{{/lignes}}']): t.rows[1].cells[i].text = v
d.add_paragraph('Total : {{total}} €')
d.add_paragraph('Cordialement, {{expediteur}}')
d.save('www/template.docx')
